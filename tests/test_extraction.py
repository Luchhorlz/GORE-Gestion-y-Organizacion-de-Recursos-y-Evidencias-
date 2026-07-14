import tempfile
import unittest
from pathlib import Path

from backend.extraction import chunk_text, extract_document


class ExtractionTests(unittest.TestCase):
    def test_text_is_normalized_and_chunked(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "prueba.txt"
            path.write_text("Primera línea.\n\n\nSegunda   línea.", encoding="utf-8")
            sections = extract_document(path, "text/plain")
            self.assertEqual(sections[0].text, "Primera línea.\n\nSegunda línea.")
            self.assertEqual(chunk_text(sections[0].text, 20, 4)[0], "Primera línea.")

    def test_word_and_excel_keep_document_sections(self):
        from docx import Document
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            word_path = root / "prueba.docx"
            word = Document()
            word.add_paragraph("Contenido del escrito judicial")
            word.save(word_path)
            self.assertIn("escrito judicial", extract_document(word_path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")[0].text)

            excel_path = root / "prueba.xlsx"
            workbook = Workbook()
            workbook.active.title = "Cronología"
            workbook.active.append(["Fecha", "Acontecimiento"])
            workbook.active.append(["01/07/2026", "Hito principal"])
            workbook.save(excel_path)
            sections = extract_document(excel_path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.assertEqual(sections[0].section_label, "Cronología")
            self.assertIn("Hito principal", sections[0].text)

    def test_pdf_preserves_page_number(self):
        from reportlab.pdfgen.canvas import Canvas

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "prueba.pdf"
            canvas = Canvas(str(path))
            canvas.drawString(72, 760, "Primera pagina del expediente")
            canvas.showPage()
            canvas.drawString(72, 760, "Segunda pagina con evidencia")
            canvas.save()
            sections = extract_document(path, "application/pdf")
            self.assertEqual([section.section_label for section in sections], ["Página 1", "Página 2"])
            self.assertIn("Segunda pagina", sections[1].text)

    def test_image_ocr_runs_entirely_local(self):
        from PIL import Image, ImageDraw, ImageFont

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "captura.png"
            image = Image.new("RGB", (1200, 250), "white")
            draw = ImageDraw.Draw(image)
            font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 72)
            draw.text((40, 70), "GORE EVIDENCIA 1 JULIO", fill="black", font=font)
            image.save(path)
            section = extract_document(path, "image/png")[0]
            self.assertEqual(section.method, "rapidocr")
            self.assertIn("GORE EVIDENCIA", section.text)


if __name__ == "__main__":
    unittest.main()
